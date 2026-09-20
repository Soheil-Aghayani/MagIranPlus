import os
import unittest
from io import BytesIO
from unittest.mock import patch
from urllib.parse import parse_qs, urlparse

from docx import Document

from citation_formats import format_citation, normalize_style
from server import app
from test_parser import SAMPLE_HTML, SOURCE_URL


ARTICLE = {
    "id": "12",
    "title": "عنوان پژوهش ۱۴۰۲ English",
    "venue": "نشریهٔ آزمون",
    "volume": "4",
    "issue": "2",
    "pages": "صص 10 -20",
    "year": "۱۴۰۲",
    "type": "مقاله ژورنالی",
    "authors": "رضا خاکپور، ناصر مهردادی",
    "url": "https://www.magiran.com/paper/12/article",
}


class CitationFormatTests(unittest.TestCase):
    def test_style_aliases_and_formats(self):
        self.assertEqual(normalize_style("APA 7th"), "apa7")
        self.assertIn("(۱۴۰۲)", format_citation(ARTICLE, 1, "apa7", "ناصر مهردادی"))
        self.assertTrue(format_citation(ARTICLE, 2, "vancouver", "ناصر مهردادی").startswith("2."))
        self.assertIn("@article", format_citation(ARTICLE, 1, "bibtex", "ناصر مهردادی"))
        self.assertIn("journal", format_citation(ARTICLE, 1, "bibtex", "ناصر مهردادی"))

    def test_isolation_preserves_all_authors_by_default(self):
        full = format_citation(ARTICLE, 1, "apa7", "ناصر مهردادی")
        isolated = format_citation(ARTICLE, 1, "apa7", "ناصر مهردادی", True, "ناصر مهردادی", True)
        self.assertIn("رضا خاکپور", full)
        self.assertEqual(isolated.split(" (")[0], "ناصر مهردادی")


class ApiTests(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()

    def test_search_validation_rejects_other_domains(self):
        response = self.client.post("/api/parse-search", json={"url": "https://example.com/searchinpapers"})
        self.assertEqual(response.status_code, 400)

    def test_health_identifies_magiranplus_service(self):
        response = self.client.get(
            "/api/health",
            headers={"Origin": "https://soheil-aghayani.github.io"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.get_json(),
            {"ok": True, "service": "MagIranPlus", "fetch_route": "direct"},
        )
        self.assertEqual(response.headers["Access-Control-Allow-Origin"], "https://soheil-aghayani.github.io")

    @patch.dict(os.environ, {"MAGIRAN_EGRESS_PROXY": "https://proxy.example.test:8443"})
    def test_health_reports_administrator_proxy_without_exposing_it(self):
        response = self.client.get("/api/health")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["fetch_route"], "configured-proxy")
        self.assertNotIn("proxy.example.test", response.get_data(as_text=True))

    def test_pasted_single_page_is_explicitly_incomplete(self):
        response = self.client.post(
            "/api/parse-html",
            json={"source_url": SOURCE_URL, "html": SAMPLE_HTML},
        )
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertFalse(payload["complete"])
        self.assertEqual(payload["source"], "pasted-html")
        self.assertEqual(payload["page_status"][0]["status"], "ok")

    def test_parse_search_fetches_all_pages_and_reports_status(self):
        page_two = SAMPLE_HTML.replace('id="fa_101"', 'id="fa_201"').replace(
            "مقالهٔ اول ۱۴۰۲", "مقالهٔ سوم"
        ).replace("ردیف ۱-۲ از ۳", "ردیف ۳-۳ از ۳")

        def fake_fetch(url):
            page = parse_qs(urlparse(url).query).get("page", ["1"])[0]
            return page_two if page == "2" else SAMPLE_HTML

        with patch("server.fetch_magiran_html", side_effect=fake_fetch):
            response = self.client.post("/api/parse-search", json={"url": SOURCE_URL, "fetch_all": True})

        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertTrue(payload["complete"])
        self.assertEqual(payload["count"], 3)
        self.assertEqual(len(payload["page_status"]), 2)

    def test_parse_search_marks_failed_page_as_incomplete(self):
        def fake_fetch(url):
            if "page=2" in url:
                raise RuntimeError("page unavailable")
            return SAMPLE_HTML

        with patch("server.fetch_magiran_html", side_effect=fake_fetch):
            response = self.client.post("/api/parse-search", json={"url": SOURCE_URL, "fetch_all": True})

        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertFalse(payload["complete"])
        self.assertTrue(any(item["status"] == "error" for item in payload["page_status"]))


class WordExportTests(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()
        self.payload = {
            "profile": {"name": "ناصر مهردادی", "url": SOURCE_URL},
            "articles": [ARTICLE],
            "style": "apa7",
        }

    def test_docx_uses_b_nazanin_14_tnr_13_rtl_and_persian_text_digits(self):
        response = self.client.post("/api/export-word", json={**self.payload, "file_type": "docx"})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.content_type,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        document = Document(BytesIO(response.data))
        citation = document.paragraphs[-1]
        self.assertIn("۱۴۰۲", citation.text)
        self.assertNotIn("1402", citation.text)
        self.assertIn("B Nazanin", {run.font.name for run in citation.runs})
        self.assertIn("Times New Roman", {run.font.name for run in citation.runs})
        self.assertIn(14.0, {run.font.size.pt for run in citation.runs if run.font.size})
        self.assertIn(13.0, {run.font.size.pt for run in citation.runs if run.font.size})
        self.assertTrue(all('w:val="start"' in paragraph._p.xml for paragraph in document.paragraphs))
        self.assertTrue(all("w:bidi" in paragraph._p.xml for paragraph in document.paragraphs))

    def test_docx_can_omit_magiran_links(self):
        response = self.client.post(
            "/api/export-word",
            json={**self.payload, "include_links": False},
        )
        document = Document(BytesIO(response.data))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertNotIn("https://www.magiran.com", text)
        self.assertIn("منبع: جست‌وجوی مگ‌ایران", text)

    def test_docx_can_isolate_target_author(self):
        response = self.client.post(
            "/api/export-word",
            json={
                **self.payload,
                "target_author": "ناصر مهردادی",
                "isolate_author": True,
            },
        )
        document = Document(BytesIO(response.data))
        citation = document.paragraphs[-1].text
        self.assertIn("ناصر مهردادی", citation)
        self.assertNotIn("رضا خاکپور", citation)


if __name__ == "__main__":
    unittest.main()
