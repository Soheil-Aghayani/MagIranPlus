"""Local-first Magiran search extractor with Persian Word exports."""

from __future__ import annotations

import io
import os
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import parse_qsl, urlencode, urlparse, urlunparse
from urllib.request import ProxyHandler, Request, build_opener, urlopen

from flask import Flask, jsonify, request, send_file, send_from_directory

from citation_formats import format_citation, normalize_style, style_label
from magiran_parser import merge_articles, parse_search_html


ROOT = Path(__file__).resolve().parent
WEB_DIR = ROOT / "web"
MAGIRAN_HOSTS = {"magiran.com", "www.magiran.com"}
MAX_HTML_BYTES = 8 * 1024 * 1024
MAX_ARTICLES = 2000
MAX_PAGES = 200
PAGE_FETCH_WORKERS = 3
PAGE_FETCH_TIMEOUT = 30
USER_AGENT = "MagIranPlus/1.0 (local academic research utility)"
EGRESS_PROXY_ENV = "MAGIRAN_EGRESS_PROXY"
PERSIAN_FONT = "B Nazanin"
ENGLISH_FONT = "Times New Roman"
WORD_PERSIAN_SIZE = 14
WORD_ENGLISH_SIZE = 13
_WESTERN_DIGITS = "0123456789"
_PERSIAN_DIGITS = "۰۱۲۳۴۵۶۷۸۹"
_ARABIC_DIGITS = "٠١٢٣٤٥٦٧٨٩"
_DIGIT_TRANSLATION = str.maketrans(
    _PERSIAN_DIGITS + _ARABIC_DIGITS,
    _WESTERN_DIGITS + _WESTERN_DIGITS,
)
_PERSIAN_DIGIT_TRANSLATION = str.maketrans(
    _WESTERN_DIGITS + _ARABIC_DIGITS,
    _PERSIAN_DIGITS + _PERSIAN_DIGITS,
)
_URL_PATTERN = re.compile(r"https?://[^\s]+", re.IGNORECASE)
_PERSIAN_CHAR_PATTERN = re.compile(r"[\u0600-\u06ff\u0750-\u077f\u08a0-\u08ff\ufb50-\ufdff\ufe70-\ufeff]")
_LATIN_CHAR_PATTERN = re.compile(r"[A-Za-z]")

app = Flask(__name__, static_folder=str(WEB_DIR), static_url_path="")
app.config["MAX_CONTENT_LENGTH"] = MAX_HTML_BYTES + 256 * 1024


def allowed_origins() -> set[str]:
    configured = os.environ.get("MAGIRAN_ALLOWED_ORIGINS", "")
    origins = {origin.strip() for origin in configured.split(",") if origin.strip()}
    origins.update({
        "https://soheil-aghayani.github.io",
        "http://127.0.0.1:5000",
        "http://localhost:5000",
    })
    return origins


@app.after_request
def add_cors_headers(response):
    origin = request.headers.get("Origin", "")
    if origin in allowed_origins():
        response.headers["Access-Control-Allow-Origin"] = origin
        response.headers["Access-Control-Allow-Methods"] = "GET, POST, OPTIONS"
        response.headers["Access-Control-Allow-Headers"] = "Content-Type"
        response.headers["Access-Control-Expose-Headers"] = "Content-Disposition"
        response.headers["Vary"] = "Origin"
    return response


def json_error(message: str, status: int = 400):
    return jsonify({"ok": False, "error": message}), status


def safe_text(value: object, limit: int = 12000) -> str:
    return str(value or "").strip()[:limit]


def boolean_value(value: object, default: bool = True) -> bool:
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    return default if not normalized else normalized not in {"0", "false", "no", "off"}


def canonical_search_url(value: str) -> str:
    raw = (value or "").strip()
    parsed = urlparse(raw)
    host = (parsed.hostname or "").lower()
    if parsed.scheme not in {"http", "https"} or host not in MAGIRAN_HOSTS:
        raise ValueError("فقط لینک جست‌وجوی مگ‌ایران پذیرفته می‌شود.")
    if parsed.path.rstrip("/").lower() != "/searchinpapers":
        raise ValueError("لینک باید شبیه https://www.magiran.com/searchinpapers?... باشد.")
    if not parsed.query:
        raise ValueError("لینک جست‌وجوی مگ‌ایران باید عبارت یا فیلتر جست‌وجو داشته باشد.")
    return urlunparse(("https", "www.magiran.com", "/searchinpapers", "", parsed.query, ""))


def page_url(source_url: str, page: int) -> str:
    parsed = urlparse(source_url)
    query = [(key, value) for key, value in parse_qsl(parsed.query, keep_blank_values=True) if key.lower() != "page"]
    if page > 1:
        query.append(("page", str(page)))
    return urlunparse((parsed.scheme, parsed.netloc, parsed.path, "", urlencode(query), ""))


def decode_html(raw: bytes, content_type: str = "") -> str:
    charset_match = re.search(r"charset=([\w-]+)", content_type or "", re.IGNORECASE)
    encoding = charset_match.group(1) if charset_match else "utf-8"
    return raw.decode(encoding, errors="replace")


def configured_egress_proxy() -> str:
    """Return the administrator-configured server egress proxy, if any.

    This is intentionally an environment setting, never a request parameter.
    User-provided V2Ray links must not become server-side proxy credentials.
    """
    return os.environ.get(EGRESS_PROXY_ENV, "").strip()


def magiran_opener():
    proxy_url = configured_egress_proxy()
    if not proxy_url:
        return None

    parsed = urlparse(proxy_url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise RuntimeError(
            f"مقدار {EGRESS_PROXY_ENV} باید یک آدرس HTTP یا HTTPS معتبر باشد."
        )

    return build_opener(ProxyHandler({"http": proxy_url, "https": proxy_url}))


def fetch_magiran_html(
    source_url: str,
    *,
    timeout: int = PAGE_FETCH_TIMEOUT,
    max_bytes: int = MAX_HTML_BYTES,
) -> str:
    headers = {
        "User-Agent": USER_AGENT,
        "Accept": "text/html,application/xhtml+xml",
        "Accept-Language": "fa-IR,fa;q=0.9,en;q=0.5",
        "Accept-Encoding": "identity",
    }
    try:
        request = Request(source_url, headers=headers)
        opener = magiran_opener()
        response_context = (
            opener.open(request, timeout=timeout)
            if opener
            else urlopen(request, timeout=timeout)
        )
        with response_context as response:
            final_host = (urlparse(response.geturl()).hostname or "").lower()
            if final_host not in MAGIRAN_HOSTS:
                raise RuntimeError("مگ‌ایران به یک مقصد ناشناس هدایت کرد.")
            raw = response.read(max_bytes + 1)
            if len(raw) > max_bytes:
                raise ValueError("حجم صفحهٔ مگ‌ایران بیشتر از حد مجاز است.")
            return decode_html(raw, response.headers.get("Content-Type", ""))
    except HTTPError as error:
        if error.code in {401, 403, 429}:
            raise RuntimeError(
                "مگ‌ایران دسترسی مستقیم این درخواست را محدود کرده است؛ حالت HTML را امتحان کنید."
            ) from error
        raise RuntimeError(f"مگ‌ایران با خطای {error.code} پاسخ داد.") from error
    except (URLError, TimeoutError, OSError) as error:
        raise RuntimeError(
            "اتصال به مگ‌ایران برقرار نشد؛ می‌توانید HTML صفحه را در حالت جایگزین وارد کنید."
        ) from error


def sanitize_articles(value: object) -> list[dict[str, str]]:
    if not isinstance(value, list):
        raise ValueError("فهرست مقاله‌ها معتبر نیست.")

    fields = {
        "id", "title", "title_en", "authors", "authors_en", "venue", "volume", "issue",
        "season", "issue_text", "year", "pages_start", "pages_end", "pages", "language",
        "abstract", "type", "url",
    }
    cleaned: list[dict[str, str]] = []
    for item in value[:MAX_ARTICLES]:
        if not isinstance(item, dict):
            raise ValueError("یکی از مقاله‌های انتخاب‌شده معتبر نیست.")
        title = safe_text(item.get("title"))
        if not title:
            continue
        cleaned.append({key: safe_text(item.get(key), 12000 if key in {"title", "abstract"} else 2000) for key in fields})
    if not cleaned:
        raise ValueError("حداقل یک مقالهٔ معتبر انتخاب کنید.")
    return cleaned


def normalized_search_payload(
    source_url: str,
    first_page: dict[str, object],
    page_results: dict[int, dict[str, object]],
    page_status: list[dict[str, object]],
    complete: bool,
    *,
    page_count: int | None = None,
    discovered_page_count: int | None = None,
    limited: bool = False,
) -> dict[str, object]:
    ordered_pages = [page_results[page] for page in sorted(page_results)]
    articles = merge_articles(ordered_pages)[:MAX_ARTICLES]
    query = str(first_page.get("query") or "")
    total_count = int(first_page.get("total_count") or len(articles))
    effective_page_count = int(page_count or first_page.get("page_count") or 1)
    discovered_count = int(discovered_page_count or first_page.get("page_count") or effective_page_count)
    return {
        "ok": True,
        "profile": {
            "id": query,
            "name": query or "نتیجهٔ جست‌وجوی مگ‌ایران",
            "affil": "جست‌وجوی مگ‌ایران",
            "url": source_url,
        },
        "source_url": source_url,
        "query": query,
        "articles": articles,
        "count": len(articles),
        "total_count": total_count,
        "page_count": effective_page_count,
        "discovered_page_count": discovered_count,
        "limited": limited,
        "complete": complete,
        "page_status": sorted(page_status, key=lambda item: int(item["page"])),
    }


def fetch_search_pages(source_url: str, fetch_all: bool = True) -> dict[str, object]:
    first_html = fetch_magiran_html(source_url)
    first_page = parse_search_html(first_html, source_url)
    if not first_page["articles"]:
        raise ValueError("در این لینک جست‌وجوی مگ‌ایران مقاله‌ای پیدا نشد.")

    discovered_page_count = max(int(first_page.get("page_count") or 1), 1)
    page_count = min(discovered_page_count, MAX_PAGES)
    page_limit_hit = discovered_page_count > MAX_PAGES
    requested_page = int(first_page.get("page") or 1)
    page_results: dict[int, dict[str, object]] = {requested_page: first_page}
    page_status: list[dict[str, object]] = [{
        "page": requested_page,
        "url": source_url,
        "status": "ok",
        "count": len(first_page["articles"]),
    }]

    if not fetch_all or page_count <= 1:
        complete = (not fetch_all and discovered_page_count <= 1) or (fetch_all and not page_limit_hit)
        return normalized_search_payload(
            source_url,
            first_page,
            page_results,
            page_status,
            complete,
            page_count=page_count,
            discovered_page_count=discovered_page_count,
            limited=page_limit_hit,
        )

    urls = {page: page_url(source_url, page) for page in range(1, page_count + 1)}
    pending = [(page, url) for page, url in urls.items() if page != requested_page]
    with ThreadPoolExecutor(max_workers=PAGE_FETCH_WORKERS) as executor:
        futures = {
            executor.submit(fetch_magiran_html, url): (page, url)
            for page, url in pending
        }
        for future in as_completed(futures):
            page, url = futures[future]
            try:
                parsed = parse_search_html(future.result(), url)
                if parsed["articles"]:
                    page_results[page] = parsed
                page_status.append({
                    "page": page,
                    "url": url,
                    "status": "ok",
                    "count": len(parsed["articles"]),
                })
            except Exception as error:
                page_status.append({
                    "page": page,
                    "url": url,
                    "status": "error",
                    "count": 0,
                    "error": str(error),
                })

    complete = (
        not page_limit_hit
        and len(page_results) == page_count
        and all(item["status"] == "ok" for item in page_status)
    )
    return normalized_search_payload(
        source_url,
        first_page,
        page_results,
        page_status,
        complete,
        page_count=page_count,
        discovered_page_count=discovered_page_count,
        limited=page_limit_hit,
    )


@app.get("/")
def index():
    return send_from_directory(WEB_DIR, "index.html")


@app.get("/api/health")
def health():
    return jsonify({
        "ok": True,
        "service": "MagIranPlus",
        "fetch_route": "configured-proxy" if configured_egress_proxy() else "direct",
    })


@app.post("/api/parse-search")
def parse_search():
    body = request.get_json(silent=True) or {}
    try:
        source_url = canonical_search_url(str(body.get("url", "")))
        payload = fetch_search_pages(source_url, fetch_all=boolean_value(body.get("fetch_all"), True))
        return jsonify(payload)
    except ValueError as error:
        return json_error(str(error), 400 if "لینک" in str(error) else 422)
    except RuntimeError as error:
        return json_error(str(error), 502)


@app.post("/api/parse-html")
def parse_html():
    body = request.get_json(silent=True) or {}
    html = str(body.get("html", ""))
    if not html.strip():
        return json_error("کد HTML صفحه را وارد کنید.", 400)
    if len(html.encode("utf-8")) > MAX_HTML_BYTES:
        return json_error("حجم HTML بیشتر از حد مجاز است.", 413)

    raw_source_url = str(body.get("source_url", "")).strip()
    if raw_source_url:
        try:
            source_url = canonical_search_url(raw_source_url)
        except ValueError:
            source_url = "https://www.magiran.com/searchinpapers"
    else:
        source_url = "https://www.magiran.com/searchinpapers"

    parsed = parse_search_html(html, source_url)
    if not parsed["articles"]:
        return json_error("از این HTML مقاله‌ای پیدا نشد. View Page Source را کپی کنید.", 422)

    page_results = {int(parsed.get("page") or 1): parsed}
    status = [{
        "page": int(parsed.get("page") or 1),
        "url": source_url,
        "status": "ok",
        "count": len(parsed["articles"]),
    }]
    payload = normalized_search_payload(source_url, parsed, page_results, status, int(parsed.get("page_count") or 1) <= 1)
    payload["source"] = "pasted-html"
    return jsonify(payload)


def set_rtl_paragraph_properties(ppr) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for tag in ("w:bidi", "w:jc"):
        for element in ppr.findall(qn(tag)):
            ppr.remove(element)

    # In an RTL paragraph, logical `start` is the visual right edge. Using
    # the logical value avoids Word treating physical `right` as the wrong
    # side when mixed Persian/Latin runs are present.
    bidi = OxmlElement("w:bidi")
    ppr.insert_element_before(
        bidi,
        "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
        "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
        "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
        "w:pPrChange",
    )
    justification = OxmlElement("w:jc")
    justification.set(qn("w:val"), "start")
    ppr.insert_element_before(
        justification,
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
        "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
        "w:pPrChange",
    )


def add_bidi(paragraph) -> None:
    set_rtl_paragraph_properties(paragraph._p.get_or_add_pPr())


def set_document_rtl_defaults(document) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    styles = document.styles.element
    doc_defaults = styles.find(qn("w:docDefaults"))
    if doc_defaults is None:
        doc_defaults = OxmlElement("w:docDefaults")
        styles.insert(0, doc_defaults)
    paragraph_defaults = doc_defaults.find(qn("w:pPrDefault"))
    if paragraph_defaults is None:
        paragraph_defaults = OxmlElement("w:pPrDefault")
        doc_defaults.append(paragraph_defaults)
    default_ppr = paragraph_defaults.find(qn("w:pPr"))
    if default_ppr is None:
        default_ppr = OxmlElement("w:pPr")
        paragraph_defaults.append(default_ppr)
    set_rtl_paragraph_properties(default_ppr)

    normal = document.styles["Normal"]
    set_rtl_paragraph_properties(normal._element.get_or_add_pPr())

    section_ppr = document.sections[0]._sectPr
    section_bidi = section_ppr.find(qn("w:bidi"))
    if section_bidi is None:
        section_ppr.insert(0, OxmlElement("w:bidi"))


def _script_chunks(value: str):
    current_script = "fa"
    buffer: list[str] = []

    def flush():
        if not buffer:
            return None
        chunk = "".join(buffer)
        buffer.clear()
        return chunk, current_script

    for character in value:
        if character in "\r\n":
            pending = flush()
            if pending:
                yield pending
            yield "\n", "fa"
            current_script = "fa"
            continue

        if _PERSIAN_CHAR_PATTERN.search(character):
            detected_script = "fa"
        elif _LATIN_CHAR_PATTERN.search(character) or character in _WESTERN_DIGITS:
            detected_script = "en" if character not in _WESTERN_DIGITS else current_script
        elif character in _PERSIAN_DIGITS or character in _ARABIC_DIGITS:
            detected_script = current_script
        else:
            detected_script = current_script

        if buffer and detected_script != current_script:
            pending = flush()
            if pending:
                yield pending
        current_script = detected_script
        buffer.append(character)

    pending = flush()
    if pending:
        yield pending


def word_text_chunks(value: object):
    text = str(value if value is not None else "")
    cursor = 0
    for match in _URL_PATTERN.finditer(text):
        before = text[cursor:match.start()]
        if before:
            for chunk, script in _script_chunks(before):
                translation = _PERSIAN_DIGIT_TRANSLATION if script == "fa" else _DIGIT_TRANSLATION
                yield chunk.translate(translation), script, False
        yield match.group(0).translate(_DIGIT_TRANSLATION), "en", True
        cursor = match.end()
    remainder = text[cursor:]
    if remainder:
        for chunk, script in _script_chunks(remainder):
            translation = _PERSIAN_DIGIT_TRANSLATION if script == "fa" else _DIGIT_TRANSLATION
            yield chunk.translate(translation), script, False


def set_run_font(run, size: float | None = None, bold: bool | None = None, font_name: str = PERSIAN_FONT, rtl: bool = True) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    run.font.name = font_name
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{slot}"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rtl_tag = rpr.find(qn("w:rtl"))
    if rtl and rtl_tag is None:
        rpr.append(OxmlElement("w:rtl"))
    elif not rtl and rtl_tag is not None:
        rpr.remove(rtl_tag)


def add_word_text(paragraph, value: object, size: float = WORD_PERSIAN_SIZE, bold: bool = False) -> None:
    for chunk, script, _is_url in word_text_chunks(value):
        if chunk == "\n":
            if paragraph.runs:
                paragraph.runs[-1].add_break()
            else:
                paragraph.add_run().add_break()
            continue
        run = paragraph.add_run(chunk)
        set_run_font(
            run,
            size=size if script == "fa" else WORD_ENGLISH_SIZE if size == WORD_PERSIAN_SIZE else max(size - 1, 1),
            bold=bold,
            font_name=PERSIAN_FONT if script == "fa" else ENGLISH_FONT,
            rtl=script == "fa",
        )


def build_docx(payload: dict[str, object]) -> io.BytesIO:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    document = Document()
    set_document_rtl_defaults(document)
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = PERSIAN_FONT
    normal.font.size = Pt(WORD_PERSIAN_SIZE)
    normal_rpr = normal._element.get_or_add_rPr()
    normal_rfonts = normal_rpr.rFonts
    if normal_rfonts is None:
        normal_rfonts = OxmlElement("w:rFonts")
        normal_rpr.insert(0, normal_rfonts)
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        normal_rfonts.set(qn(f"w:{slot}"), PERSIAN_FONT)

    profile = payload.get("profile") or {}
    query_name = str(profile.get("name") or "نتیجهٔ جست‌وجوی مگ‌ایران")
    source_url = str(profile.get("url") or "")
    articles = list(payload.get("articles") or [])[:MAX_ARTICLES]
    selected_style = normalize_style(payload.get("style"))
    selected_style_label = style_label(selected_style)
    include_links = boolean_value(payload.get("include_links"), True)
    target_author = safe_text(payload.get("target_author"), 500)
    isolate_author = boolean_value(payload.get("isolate_author"), False)
    citations = [
        format_citation(article, index, selected_style, query_name, include_links, target_author, isolate_author)
        for index, article in enumerate(articles, start=1)
    ]

    def paragraph_with_text(text: object, bold: bool = False):
        paragraph = document.add_paragraph()
        add_bidi(paragraph)
        add_word_text(paragraph, text, size=WORD_PERSIAN_SIZE, bold=bold)
        return paragraph

    paragraph_with_text("فهرست منابع مگ‌ایران", bold=True)
    paragraph_with_text(query_name, bold=True)
    info = document.add_paragraph()
    add_bidi(info)
    add_word_text(info, "تعداد مقالات انتخاب‌شده: ", bold=True)
    add_word_text(info, len(articles))
    add_word_text(info, "  |  سبک ارجاع: ", bold=True)
    add_word_text(info, selected_style_label)
    if include_links:
        add_word_text(info, "  |  منبع: ", bold=True)
        add_word_text(info, source_url)
    else:
        add_word_text(info, "  |  منبع: جست‌وجوی مگ‌ایران")

    note = (
        "این فایل بر اساس اطلاعات نمایه‌شده در نتایج جست‌وجوی مگ‌ایران ساخته شده است. "
        "اطلاعات کتاب‌شناختی را پیش از ارسال نهایی بررسی کنید."
    )
    paragraph_with_text(note)

    for citation in citations:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(7)
        add_bidi(paragraph)
        add_word_text(paragraph, citation)

    output = io.BytesIO()
    document.core_properties.title = f"منابع مگ‌ایران - {query_name}"
    document.core_properties.subject = f"فهرست منابع با سبک {selected_style_label}"
    document.save(output)
    output.seek(0)
    return output


@app.post("/api/export-word")
@app.post("/api/export-docx")
def export_word():
    body = request.get_json(silent=True) or {}
    try:
        articles = sanitize_articles(body.get("articles"))
        profile = body.get("profile") if isinstance(body.get("profile"), dict) else {}
        payload = {
            "profile": {
                "name": safe_text(profile.get("name"), 500),
                "url": safe_text(profile.get("url"), 1000),
            },
            "articles": articles,
            "style": normalize_style(body.get("style")),
            "include_links": boolean_value(body.get("include_links"), True),
            "target_author": safe_text(body.get("target_author"), 500),
            "isolate_author": boolean_value(body.get("isolate_author"), False),
        }
        output = build_docx(payload)
    except ValueError as error:
        return json_error(str(error), 400)
    except ImportError:
        return json_error("کتابخانهٔ ساخت فایل Word نصب نیست. requirements.txt را نصب کنید.", 503)
    except Exception:
        app.logger.exception("Word export failed")
        return json_error("ساخت فایل Word انجام نشد.", 500)

    filename = f"magiran-references-{normalize_style(body.get('style'))}-{datetime.now(timezone.utc):%Y%m%d}.docx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@app.errorhandler(413)
def request_too_large(_error):
    return json_error("حجم درخواست بیشتر از حد مجاز است.", 413)


if __name__ == "__main__":
    host = os.environ.get("HOST", "127.0.0.1")
    port = int(os.environ.get("PORT", "5000"))
    app.run(host=host, port=port, debug=False)
